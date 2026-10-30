# изменено 2026-10-17 05:30
"""Карта склада: какой товар на каком месте по кругу (по часовой стрелке от двери).

    # 1. сделать бланк для заполнения на складе
    python3 cake-warehouse/db/import_places.py --template karta-sklada.xlsx
    # 2. загрузить заполненный бланк → places.csv (его читает pick_sheet.py)
    python3 cake-warehouse/db/import_places.py karta-sklada.xlsx
"""
from __future__ import annotations

import argparse
import csv
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(ROOT / "api"))
from bots.cake import CATALOG, CATS  # noqa: E402

PLACES = HERE / "places.csv"
HEAD = ["Место", "Категория", "Наименование", "Вес, г"]


def template(path: Path) -> None:
    F = "Arial"
    thin = Side(style="thin", color="B0B0B0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wb = Workbook()
    ws = wb.active
    ws.title = "Карта склада"
    ws["A1"] = "Карта склада — места товаров по кругу"
    ws["A1"].font = Font(name=F, size=14, bold=True)
    ws["A2"] = ("Встаньте у двери и идите по часовой стрелке. Первое место — 1, следующее — 2 и так далее. "
                "Впишите номер места в жёлтую колонку для каждого товара. Несколько товаров на одном месте — один номер. "
                "Товар не стоит на складе — оставьте пусто.")
    ws["A2"].font = Font(name=F, size=10, italic=True)
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:D2")
    ws.row_dimensions[2].height = 58
    for c, h in enumerate(HEAD, 1):
        x = ws.cell(4, c, h)
        x.font = Font(name=F, bold=True, color="FFFFFF")
        x.fill = PatternFill("solid", fgColor="444444")
        x.border = border
        x.alignment = Alignment(horizontal="center")
    yellow = PatternFill("solid", fgColor="FFF2A8")
    r = 4
    for cat in CATS:
        for name, w in CATALOG[cat]:
            r += 1
            for c, v in enumerate([None, cat, name, w], 1):
                x = ws.cell(r, c, v)
                x.font = Font(name=F, size=11, bold=c == 1)
                x.border = border
            ws.cell(r, 1).fill = yellow
            ws.cell(r, 1).alignment = Alignment(horizontal="center")
    for col, wd in zip("ABCD", [9, 18, 40, 9]):
        ws.column_dimensions[col].width = wd
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:D{r}"
    ws.print_title_rows = "4:4"
    wb.save(path)
    print(f"Бланк: {path} ({r - 4} товаров)")


def load(path: Path) -> None:
    ws = load_workbook(path, data_only=True).active
    known = {(c, n, w) for c in CATS for n, w in CATALOG[c]}
    rows, errors, empty = [], [], []
    for i, row in enumerate(ws.iter_rows(min_row=5, values_only=True), 5):
        place, cat, name, w = (list(row) + [None] * 4)[:4]
        if not name:
            continue
        try:
            w = int(w)
        except (TypeError, ValueError):
            errors.append(f"строка {i}: вес «{w}»")
            continue
        if (cat, name, w) not in known:
            errors.append(f"строка {i}: нет в каталоге — {cat} «{name}» {w} г")
            continue
        if place in (None, ""):
            empty.append(f"{name} {w} г")
            rows.append(("", cat, name, w))
            continue
        try:
            p = int(place)
            assert p > 0
        except (TypeError, ValueError, AssertionError):
            errors.append(f"строка {i}: место «{place}» — нужно целое число больше 0")
            continue
        rows.append((p, cat, name, w))
    if errors:
        sys.exit("Не загружено, исправьте бланк:\n  " + "\n  ".join(errors))
    with PLACES.open("w", encoding="utf-8", newline="") as f:
        wr = csv.writer(f, delimiter=";")
        wr.writerow(["place", "category", "name", "weight_g"])
        wr.writerows(sorted(rows, key=lambda r: (r[0] == "", r[0] or 0, r[1], r[2])))
    placed = len(rows) - len(empty)
    print(f"{PLACES.name}: {placed} товаров с местом, мест {len({r[0] for r in rows if r[0] != ''})}")
    if empty:
        print(f"Без места ({len(empty)}) — будут в конце листа: " + ", ".join(empty[:15]) + (" …" if len(empty) > 15 else ""))


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx", nargs="?", help="заполненный бланк")
    ap.add_argument("--template", help="сделать пустой бланк сюда")
    a = ap.parse_args()
    if a.template:
        template(Path(a.template))
    elif a.xlsx:
        load(Path(a.xlsx))
    else:
        ap.error("нужен файл бланка или --template")


if __name__ == "__main__":
    main()
